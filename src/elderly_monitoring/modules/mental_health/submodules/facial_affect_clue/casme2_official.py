from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


SCHEMA_VERSION = "casme2_official_manifest_v1"
AUDIT_SCHEMA_VERSION = "casme2_official_audit_v1"
IMAGE_EXTENSIONS = {".bmp", ".jpeg", ".jpg", ".png", ".tif", ".tiff"}
PUBLICATION_RESTRICTED_SUBJECTS = {"sub12", "sub22"}
FRAME_NUMBER = re.compile(r"(\d+)$")


class Casme2OfficialAuditError(ValueError):
    """Raised when official CASME II sources cannot be mapped safely."""


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _canonical_json_bytes(value: Any) -> bytes:
    return json.dumps(
        value,
        ensure_ascii=True,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")


def _collection_sha256(rows: Iterable[dict[str, Any]]) -> str:
    digest = hashlib.sha256()
    for row in rows:
        digest.update(_canonical_json_bytes(row))
        digest.update(b"\n")
    return digest.hexdigest()


def _parse_subject(value: Any, source: str) -> tuple[str, str]:
    raw = str(value).strip()
    if not raw.isdigit():
        raise Casme2OfficialAuditError(f"Invalid subject {value!r} at {source}")
    number = int(raw)
    if number <= 0:
        raise Casme2OfficialAuditError(f"Invalid subject {value!r} at {source}")
    source_subject_id = f"{number:02d}"
    return source_subject_id, f"sub{source_subject_id}"


def _parse_required_frame(value: Any, field: str, source: str) -> int:
    if not isinstance(value, (int, float)) or int(value) != value or int(value) <= 0:
        raise Casme2OfficialAuditError(
            f"Invalid {field} frame {value!r} at {source}"
        )
    return int(value)


def _parse_optional_apex(value: Any, source: str) -> int | None:
    if value in (None, "", "/"):
        return None
    return _parse_required_frame(value, "apex", source)


def _load_coding_rows(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["Sheet1"]
    expected_headers = {
        1: "Subject",
        2: "Filename",
        4: "OnsetFrame",
        5: "ApexFrame",
        6: "OffsetFrame",
        8: "Action Units",
        9: "Estimated Emotion",
    }
    headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    for column, expected in expected_headers.items():
        if len(headers) < column or headers[column - 1] != expected:
            raise Casme2OfficialAuditError(
                f"Unexpected coding header at column {column}: "
                f"{headers[column - 1] if len(headers) >= column else None!r}"
            )

    rows: list[dict[str, Any]] = []
    for row_number, values in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        if values[0] is None:
            continue
        source = f"{path.name}:Sheet1:{row_number}"
        source_subject_id, subject_id = _parse_subject(values[0], source)
        sequence_id = str(values[1]).strip()
        if not sequence_id:
            raise Casme2OfficialAuditError(f"Empty sequence id at {source}")
        onset = _parse_required_frame(values[3], "onset", source)
        apex = _parse_optional_apex(values[4], source)
        offset = _parse_required_frame(values[5], "offset", source)
        if onset >= offset or (apex is not None and not onset <= apex <= offset):
            raise Casme2OfficialAuditError(
                f"Invalid onset/apex/offset order at {source}: "
                f"{onset}/{apex}/{offset}"
            )
        rows.append(
            {
                "source_subject_id": source_subject_id,
                "subject_id": subject_id,
                "sequence_id": sequence_id,
                "onset_frame": onset,
                "apex_frame": apex,
                "offset_frame": offset,
                "action_units": str(values[7]).strip() if values[7] is not None else None,
                "estimated_emotion": (
                    str(values[8]).strip().lower() if values[8] is not None else None
                ),
                "source_record": source,
            }
        )
    return rows


def _load_objective_classes(path: Path) -> dict[tuple[str, str], int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["Sheet1"]
    headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    if headers[:3] != ["Subject", "Filename", "Objective Class"]:
        raise Casme2OfficialAuditError(f"Unexpected objective-class headers: {headers[:3]}")

    classes: dict[tuple[str, str], int] = {}
    for row_number, values in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        if values[0] is None:
            continue
        source = f"{path.name}:Sheet1:{row_number}"
        _, subject_id = _parse_subject(values[0], source)
        sequence_id = str(values[1]).strip()
        key = (subject_id, sequence_id)
        if key in classes:
            raise Casme2OfficialAuditError(f"Duplicate objective-class key: {key}")
        value = values[2]
        if not isinstance(value, (int, float)) or int(value) != value:
            raise Casme2OfficialAuditError(
                f"Invalid objective class {value!r} at {source}"
            )
        objective_class = int(value)
        if objective_class not in range(1, 8):
            raise Casme2OfficialAuditError(
                f"Objective class outside 1-7 at {source}: {objective_class}"
            )
        classes[key] = objective_class
    return classes


def _frame_number(path: Path) -> int | None:
    match = FRAME_NUMBER.search(path.stem)
    return int(match.group(1)) if match else None


def _inventory_sequence(
    directory: Path,
    *,
    root: Path,
    hash_media: bool,
) -> dict[str, Any]:
    if not directory.is_dir():
        return {
            "path": directory.resolve().as_posix(),
            "available": False,
        "frame_count": 0,
        "first_frame": None,
        "last_frame": None,
        "frame_numbers": [],
        "collection_sha256": None,
        }

    frames = sorted(
        (
            path
            for path in directory.iterdir()
            if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS
        ),
        key=lambda path: (_frame_number(path) is None, _frame_number(path), path.name),
    )
    hash_rows = []
    for frame in frames:
        row = {
            "path": frame.relative_to(root).as_posix(),
            "size": frame.stat().st_size,
        }
        if hash_media:
            row["sha256"] = _sha256_file(frame)
        hash_rows.append(row)
    numbers = [number for frame in frames if (number := _frame_number(frame)) is not None]
    return {
        "path": directory.resolve().as_posix(),
        "available": True,
        "frame_count": len(frames),
        "first_frame": min(numbers) if numbers else None,
        "last_frame": max(numbers) if numbers else None,
        "frame_numbers": numbers,
        "collection_sha256": _collection_sha256(hash_rows),
    }


def _sequence_directories(root: Path) -> set[tuple[str, str]]:
    if not root.is_dir():
        return set()
    return {
        (subject_dir.name, sequence_dir.name)
        for subject_dir in root.iterdir()
        if subject_dir.is_dir()
        for sequence_dir in subject_dir.iterdir()
        if sequence_dir.is_dir()
    }


def audit_casme2_official(
    *,
    dataset_root: Path,
    agreement_path: Path,
    output_dir: Path,
    access_basis: str = "user_confirmed_official_application_and_download",
    hash_media: bool = True,
) -> dict[str, Any]:
    dataset_root = dataset_root.resolve()
    agreement_path = agreement_path.resolve()
    output_dir = output_dir.resolve()
    coding_path = dataset_root / "CASME2-coding-20140508.xlsx"
    objective_path = dataset_root / "CASME2-ObjectiveClasses.xlsx"
    readme_path = dataset_root / "readme.pdf"
    note_path = dataset_root / "note.txt"
    raw_root = dataset_root / "CASME2_RAW" / "CASME2-RAW"
    selected_root = dataset_root / "CASME2_RAW_selected"
    cropped_root = dataset_root / "Cropped"

    required_files = [coding_path, objective_path, readme_path, note_path, agreement_path]
    missing = [path for path in required_files if not path.is_file()]
    required_directories = [raw_root, selected_root, cropped_root]
    missing.extend(path for path in required_directories if not path.is_dir())
    if missing:
        raise FileNotFoundError(f"Missing CASME II official sources: {missing}")

    coding_rows = _load_coding_rows(coding_path)
    objective_classes = _load_objective_classes(objective_path)
    coding_keys = [(row["subject_id"], row["sequence_id"]) for row in coding_rows]
    if len(coding_keys) != len(set(coding_keys)):
        duplicates = sorted(key for key, count in Counter(coding_keys).items() if count > 1)
        raise Casme2OfficialAuditError(f"Duplicate coding keys: {duplicates}")
    if set(coding_keys) != set(objective_classes):
        raise Casme2OfficialAuditError(
            "Coding and objective-class workbooks contain different sample keys"
        )

    issues: list[dict[str, Any]] = []
    records: list[dict[str, Any]] = []
    for row in coding_rows:
        subject_id = row["subject_id"]
        sequence_id = row["sequence_id"]
        key = (subject_id, sequence_id)
        raw = _inventory_sequence(
            raw_root / subject_id / sequence_id,
            root=raw_root,
            hash_media=hash_media,
        )
        selected = _inventory_sequence(
            selected_root / subject_id / sequence_id,
            root=selected_root,
            hash_media=hash_media,
        )
        cropped = _inventory_sequence(
            cropped_root / subject_id / sequence_id,
            root=cropped_root,
            hash_media=hash_media,
        )
        annotation_status = "complete"
        if row["apex_frame"] is None:
            annotation_status = "official_apex_missing"
            issues.append(
                {
                    "severity": "warning",
                    "code": "official_apex_missing",
                    "sample_id": f"casme2_official__{subject_id}__{sequence_id}",
                    "detail": "Official coding workbook contains '/' instead of apex.",
                }
            )
        elif row["apex_frame"] in (row["onset_frame"], row["offset_frame"]):
            annotation_status = "official_apex_boundary"
            issues.append(
                {
                    "severity": "warning",
                    "code": "official_apex_boundary",
                    "sample_id": f"casme2_official__{subject_id}__{sequence_id}",
                    "detail": (
                        "Official apex equals onset or offset; retain the source value "
                        "but exclude it from interior onset-apex-offset flow training."
                    ),
                }
            )

        required_frames = [row["onset_frame"], row["offset_frame"]]
        if row["apex_frame"] is not None:
            required_frames.append(row["apex_frame"])
        raw_frame_range_ok = (
            raw["available"]
            and raw["first_frame"] is not None
            and raw["last_frame"] is not None
            and all(frame in raw["frame_numbers"] for frame in required_frames)
        )
        if not raw_frame_range_ok:
            issues.append(
                {
                    "severity": "error",
                    "code": "raw_required_frame_missing",
                    "sample_id": f"casme2_official__{subject_id}__{sequence_id}",
                    "detail": f"Required frames {required_frames} are unavailable in RAW.",
                }
            )

        selected_complete = (
            selected["available"]
            and selected["first_frame"] == row["onset_frame"]
            and selected["last_frame"] == row["offset_frame"]
        )
        cropped_complete = (
            cropped["available"]
            and cropped["first_frame"] == row["onset_frame"]
            and cropped["last_frame"] == row["offset_frame"]
        )
        if not selected_complete or not cropped_complete:
            issues.append(
                {
                    "severity": "warning",
                    "code": "derived_sequence_boundary_mismatch",
                    "sample_id": f"casme2_official__{subject_id}__{sequence_id}",
                    "detail": {
                        "expected": [row["onset_frame"], row["offset_frame"]],
                        "selected": [selected["first_frame"], selected["last_frame"]],
                        "cropped": [cropped["first_frame"], cropped["last_frame"]],
                    },
                }
            )

        records.append(
            {
                "schema_version": SCHEMA_VERSION,
                "sample_id": f"casme2_official__{subject_id}__{sequence_id}",
                **row,
                "objective_class": objective_classes[key],
                "annotation_source": "CASME2-coding-20140508.xlsx",
                "annotation_status": annotation_status,
                "raw": raw,
                "selected": selected,
                "cropped": cropped,
                "raw_required_frames_available": raw_frame_range_ok,
                "selected_boundary_complete": selected_complete,
                "cropped_boundary_complete": cropped_complete,
                "flow_ready_with_official_apex": raw_frame_range_ok
                and row["apex_frame"] is not None
                and row["onset_frame"] < row["apex_frame"] < row["offset_frame"],
                "publication_restricted": subject_id
                in PUBLICATION_RESTRICTED_SUBJECTS,
                "project_three_class_label": None,
            }
        )

    expected_keys = set(coding_keys)
    for source_name, source_root in (
        ("raw", raw_root),
        ("selected", selected_root),
        ("cropped", cropped_root),
    ):
        extras = sorted(_sequence_directories(source_root) - expected_keys)
        if extras:
            issues.append(
                {
                    "severity": "warning",
                    "code": "unannotated_sequence_directories",
                    "source": source_name,
                    "count": len(extras),
                    "sequences": [f"{subject}/{sequence}" for subject, sequence in extras],
                }
            )

    output_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = output_dir / "casme2_official_manifest_v1.jsonl"
    manifest_text = "".join(
        json.dumps(record, ensure_ascii=True, sort_keys=True) + "\n"
        for record in records
    )
    manifest_path.write_text(manifest_text, encoding="utf-8")

    source_hashes = {
        path.name: {
            "path": path.resolve().as_posix(),
            "sha256": _sha256_file(path),
            "bytes": path.stat().st_size,
        }
        for path in required_files
    }
    error_count = sum(issue["severity"] == "error" for issue in issues)
    warning_count = sum(issue["severity"] == "warning" for issue in issues)
    audit = {
        "schema_version": AUDIT_SCHEMA_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "status": "fail"
        if error_count
        else ("pass_with_source_warnings" if warning_count else "pass"),
        "dataset_root": dataset_root.as_posix(),
        "hash_media": hash_media,
        "manifest": {
            "path": manifest_path.as_posix(),
            "row_count": len(records),
            "sha256": _sha256_file(manifest_path),
        },
        "inventory": {
            "subjects": len({record["subject_id"] for record in records}),
            "samples": len(records),
            "raw_frames": sum(record["raw"]["frame_count"] for record in records),
            "selected_frames": sum(
                record["selected"]["frame_count"] for record in records
            ),
            "cropped_frames": sum(
                record["cropped"]["frame_count"] for record in records
            ),
            "complete_official_apex": sum(
                record["apex_frame"] is not None for record in records
            ),
            "interior_official_apex": sum(
                record["apex_frame"] is not None
                and record["onset_frame"] < record["apex_frame"] < record["offset_frame"]
                for record in records
            ),
            "flow_ready_with_official_apex": sum(
                record["flow_ready_with_official_apex"] for record in records
            ),
            "publication_restricted_samples": sum(
                record["publication_restricted"] for record in records
            ),
            "emotion_counts": dict(
                sorted(Counter(record["estimated_emotion"] for record in records).items())
            ),
            "objective_class_counts": {
                str(key): value
                for key, value in sorted(
                    Counter(record["objective_class"] for record in records).items()
                )
            },
        },
        "license": {
            "access_basis": access_basis,
            "agreement_template_archived_outside_repository": True,
            "signed_or_approval_evidence_archived": False,
            "research_only": True,
            "redistribution_allowed": False,
            "publication_requires_copyright_notice": "Copyright Xiaolan Fu",
            "publication_restricted_subjects": sorted(
                PUBLICATION_RESTRICTED_SUBJECTS
            ),
            "repository_contains_restricted_media": False,
        },
        "source_files": source_hashes,
        "issues": issues,
        "issue_count": len(issues),
        "error_count": error_count,
        "warning_count": warning_count,
        "downstream_guards": {
            "use_raw_as_canonical_frames": True,
            "exclude_unannotated_raw_sequences": True,
            "missing_official_apex_requires_explicit_estimation": True,
            "three_class_mapping_not_assigned_by_data_audit": True,
            "do_not_commit_or_redistribute_media": True,
        },
    }
    audit_path = output_dir / "casme2_official_audit_v1.json"
    audit_path.write_text(
        json.dumps(audit, ensure_ascii=True, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    audit["audit_path"] = audit_path.as_posix()
    audit["audit_sha256"] = _sha256_file(audit_path)
    return audit
