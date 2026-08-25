from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


REPORT = Path("docs/deliverables/跌倒风险模块研究报告-按模板初稿.docx")
BODY_CN = "宋体"
HEADING_CN = "黑体"
LATIN = "Times New Roman"
MATH = "Cambria Math"


def set_font(target, cn, latin=LATIN):
    target.font.name = latin
    rpr = target._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key, value in (("ascii", latin), ("hAnsi", latin), ("eastAsia", cn), ("cs", latin), ("hint", "eastAsia")):
        rfonts.set(qn("w:" + key), value)


def is_math(run):
    rpr = run._element.rPr
    rfonts = rpr.rFonts if rpr is not None else None
    if rfonts is None:
        return False
    return MATH in {rfonts.get(qn("w:ascii")), rfonts.get(qn("w:hAnsi"))}


def paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


doc = Document(REPORT)
heading_names = {"Title", "Heading 1", "Heading 2", "Heading 3"}
for name in ("Normal", "Body Text", "List Bullet", "List Number", "Caption"):
    try:
        set_font(doc.styles[name], BODY_CN)
        if name in {"Normal", "Body Text"}:
            doc.styles[name].font.size = Pt(12)
    except KeyError:
        pass
for name in heading_names:
    try:
        set_font(doc.styles[name], HEADING_CN)
    except KeyError:
        pass

for paragraph in paragraphs(doc):
    cn = HEADING_CN if paragraph.style.name in heading_names else BODY_CN
    for run in paragraph.runs:
        set_font(run, MATH, MATH) if is_math(run) else set_font(run, cn)

styles = doc.styles.element
defaults = styles.find(qn("w:docDefaults"))
if defaults is None:
    defaults = OxmlElement("w:docDefaults")
    styles.insert(0, defaults)
rpr_default = defaults.find(qn("w:rPrDefault"))
if rpr_default is None:
    rpr_default = OxmlElement("w:rPrDefault")
    defaults.append(rpr_default)
rpr = rpr_default.find(qn("w:rPr"))
if rpr is None:
    rpr = OxmlElement("w:rPr")
    rpr_default.append(rpr)
rfonts = rpr.find(qn("w:rFonts"))
if rfonts is None:
    rfonts = OxmlElement("w:rFonts")
    rpr.append(rfonts)
for key, value in (("ascii", LATIN), ("hAnsi", LATIN), ("eastAsia", BODY_CN), ("cs", LATIN), ("hint", "eastAsia")):
    rfonts.set(qn("w:" + key), value)

doc.save(REPORT)
print("standardized", REPORT)
